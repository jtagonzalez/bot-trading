import sys as _sys, os as _os
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
"""
Lista, EN ORDEN DE EJECUCIÓN, cada señal del DeMarker (reentrada tras >=N
velas fuera de la banda) en los 12 activos (10 forex + BTC/ETH), usando todo
el historial de 1 minuto disponible (~51 días hábiles forex, ~61 días
cripto 24x7) — "los últimos 60 días" que pidió el usuario.

Columnas: fecha/hora UTC, activo, mercado, sentido, resultado, monto,
ganancia, saldo acumulado (si se operara $1 fijo, sin martingala, en TODAS
las señales que aparecen, forex incluido, para ver el efecto de mezclar).

Exporta operaciones_demarker.xlsx

    py listar_operaciones_demarker.py --min-velas 18
"""
import argparse
import sqlite3

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from estrategia_gold4_bb import CRIPTO, FOREX, PAYOUT, cargar_velas, indicadores
from estrategias_propias import demarker_reentrada


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--min-velas", type=int, default=18)
    args = ap.parse_args()

    c = sqlite3.connect("velas.db")
    filas = []
    for a in FOREX + CRIPTO:
        cripto = a in CRIPTO
        v = indicadores(cargar_velas(c, a, operable_forex=not cripto))
        compra, venta = demarker_reentrada(v, min_velas=args.min_velas)
        ts, cl = v.ts.values, v.close.values
        ok = v.operable.values
        n = len(v)
        for i in np.flatnonzero((compra.values | venta.values) & ok):
            if i + 1 >= n or ts[i + 1] - ts[i] != 60:
                continue
            sentido = 1 if compra.values[i] else -1
            mov = np.sign(cl[i + 1] - cl[i])
            r = mov * sentido
            resultado = "GANADA" if r > 0 else ("PERDIDA" if r < 0 else "EMPATE")
            ganancia = PAYOUT if r > 0 else (-1.0 if r < 0 else 0.0)
            dt_i = v["dt"].iloc[i]
            filas.append(dict(
                fecha_hora_utc=dt_i, fecha=dt_i.date(), hora_utc=dt_i.strftime("%H:%M"),
                activo=a, mercado="cripto" if cripto else "forex",
                sentido="COMPRA" if sentido > 0 else "VENTA",
                precio_entrada=cl[i], precio_salida=cl[i + 1],
                resultado=resultado, monto=1.0, ganancia=ganancia))

    df = pd.DataFrame(filas).sort_values("fecha_hora_utc").reset_index(drop=True)
    cols = ["fecha", "hora_utc", "fecha_hora_utc", "activo", "mercado", "sentido",
            "precio_entrada", "precio_salida", "resultado", "monto", "ganancia"]
    df = df[cols]
    df["fecha_hora_utc"] = df["fecha_hora_utc"].dt.tz_localize(None)
    df.index += 1
    df.index.name = "n_operacion"
    df["saldo_acumulado_todos"] = df.ganancia.cumsum()
    solo_cripto = df.mercado == "cripto"
    df["saldo_acumulado_solo_cripto"] = df.ganancia.where(solo_cripto, 0).cumsum()

    nombre = "operaciones_demarker.xlsx"
    hoja = f"DeMarker min{args.min_velas} velas"
    df.to_excel(nombre, sheet_name=hoja)

    # colorear la columna "resultado" segun GANADA/PERDIDA/EMPATE
    verde = PatternFill("solid", fgColor="C6EFCE")
    rojo = PatternFill("solid", fgColor="FFC7CE")
    gris = PatternFill("solid", fgColor="E7E6E6")
    texto_verde, texto_rojo, texto_gris = Font(color="006100"), Font(color="9C0006"), Font(color="666666")
    wb = load_workbook(nombre)
    ws = wb[hoja]
    col_resultado = list(df.columns).index("resultado") + 2   # +1 por el indice, +1 por 1-based
    for fila in range(2, ws.max_row + 1):
        celda = ws.cell(row=fila, column=col_resultado)
        if celda.value == "GANADA":
            celda.fill, celda.font = verde, texto_verde
        elif celda.value == "PERDIDA":
            celda.fill, celda.font = rojo, texto_rojo
        elif celda.value == "EMPATE":
            celda.fill, celda.font = gris, texto_gris
    for col in ws.columns:
        letra = col[0].column_letter
        ancho = max(len(str(c.value)) for c in col if c.value is not None)
        ws.column_dimensions[letra].width = max(ancho + 2, 10)
    wb.save(nombre)

    print(f"{len(df)} operaciones en orden de ejecución -> {nombre}\n")
    print(df.resultado.value_counts())
    print(f"\nganadas/perdidas (sin empates): {100 * (df.resultado=='GANADA').sum() / (df.resultado!='EMPATE').sum():.2f}% acierto")
    print(f"saldo final operando los 12 juntos: ${df.ganancia.sum():+.2f}")
    print(f"saldo final operando SOLO cripto:   ${df.ganancia.where(solo_cripto,0).sum():+.2f}")
    print(f"\nprimeras 5 operaciones:")
    print(df[["fecha","hora_utc","activo","mercado","sentido","resultado","ganancia"]].head(5).to_string())
    print(f"\nultimas 5 operaciones:")
    print(df[["fecha","hora_utc","activo","mercado","sentido","resultado","ganancia"]].tail(5).to_string())


if __name__ == "__main__":
    main()
